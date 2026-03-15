#!/usr/bin/env python3
"""
Mulch Order Email Parser v2
- Extracts: Name, Email, Phone, Delivery Address, Delivery Instructions,
            Bags of Mulch, Donation
- Donation: sum of all donation product lines in the order table
"""

import email
import glob
import os
import re
from email import policy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EML_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Emails")
OUTPUT  = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Mulch_Orders_2026.xlsx")

# ─────────────────────────────────────────────
# PARSER
# ─────────────────────────────────────────────

def get_text(msg):
    for part in msg.walk():
        if part.get_content_type() == 'text/plain':
            return part.get_content()
    return ""


def parse_field(text, label):
    """Extract value after a bold *Label* field."""
    pattern = rf'\*{label}\*\s*(.*?)(?=\n\s*\*|\nPayment|\nDelivery\n|\nOrder\n|$)'
    m = re.search(pattern, text, re.DOTALL | re.IGNORECASE)
    if m:
        val = m.group(1).strip()
        val = re.sub(r'Map It\s*\n?.*?(?=\n|$)', '', val, flags=re.IGNORECASE)
        val = re.sub(r'<https?://[^>]+>', '', val)
        val = re.sub(r'\n+', ' ', val).strip()
        return val if val else None
    return None


def parse_phone(text):
    """Extract phone number from the Phone (optional) field."""
    m = re.search(r'\*Phone[^*]*\*\s*(.*?)(?=\n\s*\*)', text, re.DOTALL | re.IGNORECASE)
    if not m:
        return None
    raw = m.group(1).strip()
    # Pull just the phone number
    phone_m = re.search(r'\(?\d{3}\)?[\s.\-]?\d{3}[\s.\-]?\d{4}', raw)
    return phone_m.group(0) if phone_m else (raw if raw else None)


def parse_address(text):
    """Extract delivery address block."""
    m = re.search(
        r'\*Delivery Address[^*]*\*\s*(.*?)(?=\nDelivery\n|\nPayment|\*)',
        text, re.DOTALL | re.IGNORECASE
    )
    if not m:
        return None
    lines = []
    for line in m.group(1).splitlines():
        line = line.strip()
        if not line:
            continue
        if re.match(r'Map It', line, re.IGNORECASE):
            continue
        if re.match(r'<https?://', line):
            continue
        if line.lower() == 'united states':
            continue
        lines.append(line)
    return ', '.join(lines) if lines else None


def parse_bags(text):
    """Get Bags of Mulch qty from the order table."""
    m = re.search(r'\*Bags of Mulch\*', text, re.IGNORECASE)
    if not m:
        return 0  # donation-only order
    snippet = text[m.end():m.end() + 300]
    qty_m = re.search(r'^\s*(\d+)\s+\$', snippet, re.MULTILINE)
    if qty_m:
        return int(qty_m.group(1))
    for line in snippet.splitlines():
        if re.match(r'^\d+$', line.strip()):
            return int(line.strip())
    return None


def parse_donation(text):
    """
    Sum all donation product lines in the order table.
    Donation product names: '$X Donation', 'Donation amount of your choice',
    'Enter a donation amount of your choice', 'No donation at this time'
    We extract the Price (last $ column) for each donation line.
    """
    total = 0.0

    # Find every product line that's donation-related
    for m in re.finditer(
        r'\*((?:[Nn]o\s+)?(?:\$\d+\s+)?[Dd]onation[^*]*|'
        r'(?:Enter a )?[Dd]onation amount of your choice)\*',
        text
    ):
        label = m.group(1)
        snippet = text[m.end():m.end() + 150]

        # Find "qty  $unit  $price" pattern on the next non-empty line(s)
        price_m = re.search(r'\d+\s+\$[\d.]+\s+\$([\d.]+)', snippet)
        if price_m:
            amount = float(price_m.group(1))
            total += amount
            continue

        # Fallback: last dollar amount in snippet
        all_prices = re.findall(r'\$([\d.]+)', snippet)
        if all_prices:
            total += float(all_prices[-1])

    return round(total, 2) if total > 0 else 0.0


def parse_email_file(filepath):
    with open(filepath, 'rb') as fh:
        msg = email.message_from_bytes(fh.read(), policy=policy.default)

    text = get_text(msg)
    if not text:
        return None

    name         = parse_field(text, 'Name')
    email_addr   = parse_field(text, 'Email')
    phone        = parse_phone(text)
    address      = parse_address(text)
    instructions = parse_field(text, 'Delivery instructions')
    bags         = parse_bags(text)
    donation     = parse_donation(text)

    return {
        'Name':                  name,
        'Email':                 email_addr,
        'Phone':                 phone,
        'Delivery Address':      address,
        'Delivery Instructions': instructions,
        'Bags of Mulch':         bags,
        'Donation':              donation if donation else None,
        '_filename':             os.path.basename(filepath),
    }


# ─────────────────────────────────────────────
# EXCEL WRITER
# ─────────────────────────────────────────────

HEADERS = ['Name', 'Email', 'Phone', 'Delivery Address',
           'Delivery Instructions', 'Bags of Mulch', 'Donation']

HEADER_BG  = 'FF2D6E3A'
HEADER_FG  = 'FFFFFFFF'
ALT_ROW    = 'FFF0F7F1'
BORDER_COL = 'FFCCCCCC'
DONATE_BG  = 'FFE8F5E9'  # soft green tint for donation cells


def thin_border():
    s = Side(style='thin', color=BORDER_COL)
    return Border(left=s, right=s, top=s, bottom=s)


def write_excel(records, output_path):
    wb = Workbook()

    # ── Sheet 1: All Orders ──────────────────
    ws = wb.active
    ws.title = "All Orders"

    total_orders = len(records)
    ws.merge_cells('A1:G1')
    tc = ws['A1']
    tc.value = f"AEHS Boosters — Mulch Orders 2026   ({total_orders} orders)"
    tc.font      = Font(name='Calibri', bold=True, size=14, color=HEADER_FG)
    tc.fill      = PatternFill('solid', fgColor=HEADER_BG)
    tc.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(row=2, column=col, value=h)
        c.font      = Font(name='Calibri', bold=True, color=HEADER_FG, size=10)
        c.fill      = PatternFill('solid', fgColor=HEADER_BG)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border    = thin_border()
    ws.row_dimensions[2].height = 22

    for i, rec in enumerate(records):
        row = i + 3
        for col, h in enumerate(HEADERS, 1):
            val = rec.get(h)
            c = ws.cell(row=row, column=col, value=val)
            c.font      = Font(name='Calibri', size=9)
            c.alignment = Alignment(vertical='top',
                                    wrap_text=(h in ('Delivery Address', 'Delivery Instructions')))
            c.border = thin_border()

            if h == 'Donation' and val:
                c.number_format = '$#,##0.00'
                c.fill = PatternFill('solid', fgColor=DONATE_BG)
            elif h == 'Bags of Mulch':
                c.alignment = Alignment(horizontal='center', vertical='top')
            elif i % 2 == 0:
                c.fill = PatternFill('solid', fgColor=ALT_ROW)

    # Totals row
    total_row = len(records) + 3
    ws.cell(row=total_row, column=5, value='TOTALS').font = \
        Font(name='Calibri', bold=True, size=10)
    bags_total = sum(r['Bags of Mulch'] or 0 for r in records)
    donate_total = sum(r['Donation'] or 0 for r in records)

    bc = ws.cell(row=total_row, column=6, value=bags_total)
    bc.font      = Font(name='Calibri', bold=True, size=10)
    bc.alignment = Alignment(horizontal='center')
    bc.fill      = PatternFill('solid', fgColor='FFD5E8D4')
    bc.border    = thin_border()

    dc = ws.cell(row=total_row, column=7, value=donate_total)
    dc.font          = Font(name='Calibri', bold=True, size=10)
    dc.number_format = '$#,##0.00'
    dc.fill          = PatternFill('solid', fgColor='FFD5E8D4')
    dc.border        = thin_border()

    col_widths = [22, 28, 16, 38, 44, 14, 12]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = 'A3'

    # ── Sheet 2: Summary ────────────────────
    ws2 = wb.create_sheet("Summary")
    donors    = [r for r in records if r.get('Donation')]
    no_donate = len(records) - len(donors)

    summary = [
        ("Total Orders",            len(records)),
        ("Total Bags of Mulch",     bags_total),
        ("Average Bags / Order",    round(bags_total / len(records), 1)),
        ("",                        ""),
        ("Total Donations",         f"${donate_total:,.2f}"),
        ("Orders with Donation",    len(donors)),
        ("Orders without Donation", no_donate),
    ]

    ws2.merge_cells('A1:B1')
    t = ws2['A1']
    t.value     = "Summary"
    t.font      = Font(name='Calibri', bold=True, size=14, color=HEADER_FG)
    t.fill      = PatternFill('solid', fgColor=HEADER_BG)
    t.alignment = Alignment(horizontal='center', vertical='center')
    ws2.row_dimensions[1].height = 28

    for i, (label, val) in enumerate(summary):
        r = i + 2
        lc = ws2.cell(row=r, column=1, value=label)
        lc.font = Font(name='Calibri', bold=bool(label), size=10)
        vc = ws2.cell(row=r, column=2, value=val)
        vc.font = Font(name='Calibri', size=10)

    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 18

    wb.save(output_path)
    return bags_total, donate_total, len(donors)


# ─────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────

if __name__ == '__main__':
    files = sorted(glob.glob(f'{EML_DIR}/*.eml'))
    print(f"Processing {len(files)} emails...")

    records = []
    for f in files:
        try:
            rec = parse_email_file(f)
            if rec:
                records.append(rec)
        except Exception as e:
            print(f"  ERROR: {os.path.basename(f)}: {e}")

    records.sort(key=lambda r: (r.get('Name') or '').lower())

    bags_total, donate_total, donors = write_excel(records, OUTPUT)

    print(f"\nResults:")
    print(f"  Orders parsed:    {len(records)}")
    print(f"  Total bags:       {bags_total}")
    print(f"  Total donations:  ${donate_total:,.2f}  ({donors} orders with donation)")
    print(f"  Saved: {OUTPUT}")
